"""Second-brain indexer: walks your note folders and builds a full-text index.

Incremental: a file is only re-read when its mtime or size changes.
Supported: .md .txt .tex .py .m (plain text), .pdf (per page), .docx.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Iterable

from ..config import Config
from ..db import Database, now_iso
from ..items import detect_course

CHUNK_CHARS = 1500
OVERLAP = 200


def _chunks(text: str) -> Iterable[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return
    if len(text) <= CHUNK_CHARS:
        yield text
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        # Prefer to break on a paragraph or sentence boundary.
        if end < len(text):
            for sep in ("\n\n", "\n", ". "):
                cut = text.rfind(sep, start + CHUNK_CHARS // 2, end)
                if cut != -1:
                    end = cut + len(sep)
                    break
        yield text[start:end].strip()
        if end >= len(text):
            break
        start = max(end - OVERLAP, start + 1)


def extract(path: Path) -> list[tuple[int | None, str]]:
    """Return [(page or None, text)]."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        out = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                t = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - broken pages shouldn't kill the file
                t = ""
            if t.strip():
                out.append((i, t))
        return out
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [(None, "\n".join(parts))]
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return [(None, path.read_text(encoding=enc))]
        except (UnicodeDecodeError, UnicodeError):
            continue
    return []


def guess_course(path: Path, root: Path, fixed: str, known: list[str]) -> str | None:
    if fixed:
        return fixed
    try:
        rel_parts = path.relative_to(root).parts
    except ValueError:
        rel_parts = path.parts
    for part in rel_parts:
        c = detect_course(part)
        if c:
            return c
        low = part.lower().replace("_", " ")
        for k in known:
            if k.lower() in low or k.lower().replace(" ", "") in low.replace(" ", ""):
                return k
    return None


def index_folders(db: Database, cfg: Config) -> dict[str, Any]:
    opts = cfg.notes_options
    exts = {e.lower() for e in opts["extensions"]}
    ignore = set(opts["ignore"])
    max_bytes = float(opts["max_file_mb"]) * 1024 * 1024
    known = [r["code"] for r in db.query("SELECT code FROM courses")]
    seen: set[str] = set()
    stats = {"scanned": 0, "indexed": 0, "removed": 0, "errors": 0, "missing_roots": []}

    for folder in cfg.notes_folders:
        root = cfg.resolve(folder.path)
        if not root.exists():
            stats["missing_roots"].append(str(root))
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            dirnames[:] = [d for d in dirnames if d not in ignore and not d.startswith(".")]
            for fn in filenames:
                p = Path(dirpath) / fn
                if p.suffix.lower() not in exts or fn.startswith("~$"):
                    continue
                try:
                    st = p.stat()
                except OSError:
                    continue
                if st.st_size > max_bytes:
                    continue
                key = str(p)
                seen.add(key)
                stats["scanned"] += 1
                doc = db.one("SELECT id, mtime, size FROM documents WHERE path=?", (key,))
                if doc and doc["mtime"] == st.st_mtime and doc["size"] == st.st_size:
                    continue
                course = guess_course(p, root, folder.course, known)
                try:
                    pages = extract(p)
                    err = None
                except Exception as e:  # noqa: BLE001
                    pages, err = [], f"{type(e).__name__}: {e}"[:300]
                    stats["errors"] += 1
                with db.tx() as c:
                    if doc:
                        c.execute("DELETE FROM chunks WHERE doc_id=?", (doc["id"],))
                        c.execute(
                            """UPDATE documents SET root=?, course=?, title=?, ext=?, mtime=?, size=?, indexed_at=?, error=?
                               WHERE id=?""",
                            (str(root), course, p.stem, p.suffix.lower(), st.st_mtime, st.st_size, now_iso(), err, doc["id"]),
                        )
                        doc_id = doc["id"]
                    else:
                        cur = c.execute(
                            """INSERT INTO documents(path, root, course, title, ext, mtime, size, indexed_at, error)
                               VALUES(?,?,?,?,?,?,?,?,?)""",
                            (key, str(root), course, p.stem, p.suffix.lower(), st.st_mtime, st.st_size, now_iso(), err),
                        )
                        doc_id = cur.lastrowid
                    n = 0
                    for page, text in pages:
                        for chunk in _chunks(text):
                            c.execute(
                                "INSERT INTO chunks(content, title, course, doc_id, chunk_no, page) VALUES(?,?,?,?,?,?)",
                                (chunk, p.stem, course or "", doc_id, n, page),
                            )
                            n += 1
                stats["indexed"] += 1

    # Remove documents that disappeared (only under roots we could see this run).
    live_roots = [str(cfg.resolve(f.path)) for f in cfg.notes_folders if cfg.resolve(f.path).exists()]
    for d in db.query("SELECT id, path, root FROM documents"):
        if d["path"] in seen:
            continue
        if d["root"] in live_roots or not any(d["path"].startswith(r) for r in live_roots + stats["missing_roots"]):
            with db.tx() as c:
                c.execute("DELETE FROM chunks WHERE doc_id=?", (d["id"],))
                c.execute("DELETE FROM documents WHERE id=?", (d["id"],))
            stats["removed"] += 1
    return stats


_TOKEN = re.compile(r"[\w\-\.']+", re.UNICODE)


def _fts_query(q: str, mode: str = "AND") -> str:
    tokens = [t.replace('"', "") for t in _TOKEN.findall(q) if len(t) > 1 or t.isdigit()]
    if not tokens:
        return ""
    parts = [f'"{t}"' for t in tokens]
    if mode == "AND" and len(parts) > 1:
        return " ".join(parts)
    return " OR ".join(parts)


def search(db: Database, q: str, *, course: str | None = None, limit: int = 20) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for mode in ("AND", "OR"):
        fq = _fts_query(q, mode)
        if not fq:
            return []
        sql = """SELECT c.doc_id, c.page, c.chunk_no, d.path, d.title, d.course, d.ext,
                        snippet(chunks, 0, '[[', ']]', ' … ', 24) AS snippet,
                        c.content AS content, bm25(chunks, 1.0, 4.0, 0.5) AS rank
                 FROM chunks c JOIN documents d ON d.id = c.doc_id
                 WHERE chunks MATCH ?"""
        params: list[Any] = [fq]
        if course:
            sql += " AND d.course = ?"
            params.append(course)
        sql += " ORDER BY rank LIMIT ?"
        params.append(limit)
        try:
            results = db.query(sql, tuple(params))
        except Exception:  # noqa: BLE001 - malformed FTS query
            results = []
        if results:
            break
    return results


def search_grouped(db: Database, q: str, *, course: str | None = None, limit: int = 20) -> list[dict[str, Any]]:
    """One row per document, best hit first, with up to 3 snippets."""
    hits = search(db, q, course=course, limit=limit * 3)
    grouped: dict[int, dict[str, Any]] = {}
    for h in hits:
        g = grouped.setdefault(h["doc_id"], {
            "doc_id": h["doc_id"], "path": h["path"], "title": h["title"], "course": h["course"],
            "ext": h["ext"], "snippets": [], "pages": [],
        })
        if len(g["snippets"]) < 3:
            g["snippets"].append(h["snippet"])
            if h["page"]:
                g["pages"].append(h["page"])
    return list(grouped.values())[:limit]


def doc_stats(db: Database) -> dict[str, Any]:
    r = db.one("SELECT COUNT(*) AS docs, SUM(size) AS bytes FROM documents") or {}
    by_course = db.query(
        "SELECT COALESCE(course,'—') AS course, COUNT(*) AS n FROM documents GROUP BY course ORDER BY n DESC"
    )
    return {"docs": r.get("docs") or 0, "bytes": r.get("bytes") or 0, "by_course": by_course}
